from io import BytesIO
from zipfile import ZipFile
from uuid import uuid4
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfReader
from docx import Document
from backend.policy.alternative_export import sections,router
from backend.policy.alternative_review import initialize
from tests.test_graph import service,scenario
from tests.test_regional_task import task


# /ship Step 7 pass 1 extension
# Value: protects=zero-draft PDF and Word keep frozen fact values, provenance and version-bound HTTP downloads;
# fails_when=reports omit source definitions, invent drafts, accept a wrong version or fetch fresh evidence;
# why_new=the existing test checked a PDF signature and only one Word fact ID;
# seam=none; no_claim=fake workflow and generated statistics; no live provider or policy quality acceptance.
# Pass 2: protects=observed fact status and stored reason survive both document formats;
# fails_when=status/reason lines disappear; why_new=provenance assertions omitted uncertainty; seam=none; no_claim=same fake inputs.
def test_zero_draft_report_matches_frozen_statistics_and_guidance(service,monkeypatch):
    initialize(service.store)
    service.settings.llm_max_calls=10
    run=service.start(scenario('monthly',workflow_version='regional-workflow-1',policy_context_version='regional-policy-2',task=task(dong_cd='11230001').model_dump(mode='json')))
    import json
    with service.store.connection() as db:
        body=json.loads(db.execute('SELECT body FROM policy_alternative_sets WHERE owner_id=?',(run.id,)).fetchone()[0])
    assert body['alternatives']==[]
    content=sections(body);text='\n'.join(line for _,line in content)
    assert '새 정책 대안을 작성하지 않았습니다' in text and '신청 방법:' in text
    for fact in body['evidence']['regional_profile']['facts']:
        assert fact['fact_id'] in text and str(fact['reference_year']) in text
        if fact['status']=='observed':assert str(fact['value']) in text
    app=FastAPI();app.include_router(router(service.store));client=TestClient(app)
    def forbidden(*args,**kwargs):raise AssertionError('export must use frozen evidence without tool or model calls')
    monkeypatch.setattr(service.agents,'call',forbidden)
    monkeypatch.setattr(service.policy_agent,'call_tool',forbidden)
    from backend.context.service import ContextService
    monkeypatch.setattr(ContextService,'profile',forbidden)
    path=f"/policy-proposals/{run.id}/{body['dong_cd']}/download/"
    params={'version':body['version'],'revision':0}
    for format in ('pdf','docx'):
        response=client.get(path+format,params=params)
        assert response.status_code==200 and response.headers['cache-control']=='no-store'
        exported=('\n'.join(p.extract_text() for p in PdfReader(BytesIO(response.content)).pages) if format=='pdf'
                  else '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs))
        compact=''.join(exported.split())
        assert '새정책대안을작성하지않았습니다' in compact and body['version'] in compact
        assert '검토revision0' in compact and body['dong_cd'] in compact
        for fact in body['evidence']['regional_profile']['facts']:
            for key in ('fact_id','unit','reference_year','geography_level','geography_code','geography_version',
                        'population','age_definition','sex','source_table','source_items','source_dimensions',
                        'source_hash','fetched_at','source_url'):
                assert ''.join(str(fact.get(key)).split()) in compact,(format,key,fact['fact_id'])
            assert ''.join(str(fact.get('published_at') or '미확인').split()) in compact
            assert str(fact['value'] if fact['status']=='observed' else '미확보') in compact
            assert ''.join(('관측 상태: '+fact['status']+' · 미확보 사유: '+str(fact.get('reason') or '해당 없음')).split()) in compact
    assert client.get(path+'pdf',params={**params,'version':'stale-version'}).status_code==409
    assert client.get(path+'pdf',params={**params,'revision':99}).status_code==404
    assert client.get(path+'pdf',params={**params,'alternative_id':str(uuid4())}).status_code==404
    with service.store.connection() as db:
        assert json.loads(db.execute('SELECT body FROM policy_alternative_sets WHERE owner_id=?',(run.id,)).fetchone()[0])==body


# /ship Step 7 pass 1 extension
# Value: protects=three-draft PDF and Word retain outreach text/channel, design links and each immutable review revision;
# fails_when=outreach is replaced with generic copy, linked choices disappear or old downloads acquire new approval;
# why_new=the existing case inspected sections and a PDF signature, without Word, outreach or revision content;
# seam=none; no_claim=generated frozen context and fake drafts; no message dispatch, provider or semantic quality acceptance.
# Pass 2: protects=need ownership, unmet conditions and missing-fact uncertainty remain explicit in reports;
# fails_when=need IDs/groups, missing reasons or unresolved conditions disappear; why_new=original fixtures were fully observed;
# seam=none; no_claim=renderer contracts over synthetic need reviews, not upstream decision quality.
def test_three_draft_report_keeps_design_to_fact_links(service,tmp_path,monkeypatch):
    initialize(service.store)
    from tests.test_pinpoint_policy import context
    from backend.policy.pinpoint import generate
    c=context(tmp_path);c['facts'][0].update(index_name='S1',source_period='2025-07..2025-09')
    c['need_reviews'][0]['outreach_draft']={'channel_proposal':'주민센터 게시판 검토',
        'text':'확인된 사업 안내입니다.\n신청 전 대상·기한과 <문의처>를 확인하세요.','dispatch_performed':False}
    c['design_questions']=['운영시간과 실제 접근 경로를 현장에서 확인하세요.']
    c['need_reviews'][0]['unmet_requirements']=['퇴근 후 접점의 운영 가능 시간 확인']
    c['need_reviews'].append({'need_id':'held-health-need','target_group':'F/youth','need':'health','decision':'needs_evidence',
        'reason':'material_policy_target_missing','missing_evidence':['material_policy_time_missing','unknown_contract_code','fake: 현장 확인 문장'],
        'guidance':[],'evidence_ids':['E1'],'outreach_draft':{'channel_proposal':None,'text':'미확인 조건은 담당자 검토가 필요합니다.','dispatch_performed':False}})
    from tests.test_context_analysis import replace
    from backend.tools.indicators import digest
    replace(c['regional_profile'],'total_housing',value=None,status='missing',reason='fixture_housing_not_supplied')
    c['regional_profile']['input_hash']=digest({k:v for k,v in c['regional_profile'].items() if k!='input_hash'})
    run=service.register(scenario('monthly'));scope={**c['scope'],'policy_context_version':'regional-policy-2'}
    body=generate(service.policy_designer,run.id,run.id,scope,[{'dong_cd':c['dong_cd'],'dong_nm':c['dong_nm']}],[],contexts={c['dong_cd']:c})[0]
    assert len(body['alternatives'])==3,body['fallback_reason']
    text='\n'.join(line for _,line in sections(body))
    for alternative in body['alternatives']:
        assert alternative['title'] in text
        for link in alternative['context_links']:assert link['fact_id'] in text and link['design_choice'] in text
    from backend.policy.alternative_review import update,read
    from tests.test_alternative_review import change
    approved=update(service.store,body['owner_id'],body['dong_cd'],change(body))
    app=FastAPI();app.include_router(router(service.store));client=TestClient(app)
    def forbidden(*args,**kwargs):raise AssertionError('download must not call a model')
    monkeypatch.setattr(service.agents,'call',forbidden)
    base=f"/policy-proposals/{body['owner_id']}/{body['dong_cd']}/download/"
    for revision in (0,1):
        for format in ('pdf','docx'):
            response=client.get(base+format,params={'version':body['version'],'revision':revision})
            assert response.status_code==200
            exported=('\n'.join(p.extract_text() for p in PdfReader(BytesIO(response.content)).pages) if format=='pdf'
                      else '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs))
            compact=''.join(exported.split())
            assert ('정책안검토승인' in compact)==bool(revision)
            for expected in (body['version'],f'검토 revision {revision}',c['need_reviews'][0]['outreach_draft']['text'],
                             '제안 채널: 주민센터 게시판 검토','실제 발송·게시하지 않은 초안입니다.',*c['design_questions']):
                assert ''.join(expected.split()) in compact
            for alternative in body['alternatives']:
                assert ''.join(alternative['title'].split()) in compact
                for link in alternative['context_links']:
                    for expected in link.values():assert ''.join(expected.split()) in compact
                for expected in alternative['evaluation_design'].values():assert ''.join(expected.split()) in compact
                for identifier in alternative['need_ids']:
                    assert ''.join(('다루는 필요: 여성 청년 · 관계망 · '+identifier).split()) in compact
            for expected in ('보완할 요건: 퇴근 후 접점의 운영 가능 시간 확인',
                             '공식 정책의 대상 연령·성별 조건을 확인하지 못했습니다.',
                             '미확인: 정책의 유효 시점을 확인할 자료가 부족합니다.',
                             '미확인: 세부 사유는 진행 기록에서 확인하세요.','미확인: 현장 확인 문장',
                             '제안 채널: 담당자 확인 필요','관측 상태: missing · 미확보 사유: fixture_housing_not_supplied','주택: 미확보 houses'):
                assert ''.join(expected.split()) in compact
            assert 'unknown_contract_code' not in exported and 'material_policy_time_missing' not in exported
    selected=client.get(base+'docx',params={'version':body['version'],'revision':0,'alternative_id':body['alternatives'][0]['id']})
    with ZipFile(BytesIO(selected.content)) as archive:word=archive.read('word/document.xml').decode()
    assert body['alternatives'][0]['title'] in word
    assert all(a['title'] not in word for a in body['alternatives'][1:])
    with service.store.connection() as db:
        assert read(db,body['owner_id'],body['dong_cd'],0)==body
        assert read(db,body['owner_id'],body['dong_cd'],1)==approved
        assert approved['evidence']==body['evidence']
