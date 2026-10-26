"""PDF/DOCX from one immutable review revision; no tool or model calls."""
import json
import re
from io import BytesIO
from pathlib import Path
from html import escape
from typing import Literal
from uuid import UUID
from fastapi import APIRouter,HTTPException,Query,Response
from backend.events import utcnow
from backend.policy.alternative_review import read

REASONS=json.loads((Path(__file__).resolve().parents[2]/'rules/regional_response_reasons.json').read_text())
def response_reason(reason):return REASONS.get(reason) or ('세부 사유는 진행 기록에서 확인하세요.' if re.fullmatch(r'[a-z][a-z0-9_]*',reason) else re.sub(r'^fake:\s*','',reason))

FONT=Path(__file__).resolve().parents[1]/'assets/fonts/NanumGothic-Regular.ttf'
NEEDS={'relationship':'관계망','emotional':'정서','economic':'경제','housing':'주거','employment':'고용','health':'건강'}

MODES={'llm':'모델 작성','fixture':'테스트 작성','template_fallback':'기본 구상안','mixed':'모델·구상안 혼합','not_generated':'새 대안 미작성'}
STATES={'pending':'검토 대기','approved':'정책안 검토 승인','rejected':'반려'}
EVIDENCE={'supported':'지역·참고 근거 연결','partial':'일부 근거 연결','concept_only':'근거 보완 필요'}
ASSESSMENTS={'assessed':'판정 가능','small_population_hold':'소규모 인구 보류 (0)','insufficient_evidence':'판정 근거 부족 (−1)'}
ALARMS={'baseline_unavailable':'기준선 없음 · 신호 미산출','none':'경보 없음','weak':'약한 경보','strong':'강한 경보'}

def ratio(value):return '기준선 없음' if value is None else f'{value:.3f}배'


def group_label(value):
    sex,age=value.split('/')
    return {'M':'남성','F':'여성'}.get(sex,sex)+' '+{'youth':'청년','midlife':'중년','prime':'장년','elderly':'노년'}.get(age,age)


def sections(body,alternative_id=None):
    selected=[a for a in body['alternatives'] if alternative_id is None or a['id']==alternative_id]
    if not selected and (alternative_id is not None or body.get('proposal_version')!='regional-policy-2'):raise KeyError('alternative missing')
    scope=body['scope'];result=[(1,body['dong_nm']+' 정책 제안 검토 보고서'),
        (0,'검토용 초안 · 예산 승인이나 실제 집행 문서가 아닙니다.'),
        (0,f"제안 버전 {body['version']} / 검토 revision {body['review_revision']} / 출력 {utcnow().isoformat()}"),
        (0,f"지역 {scope['sgg']} · 동 {body['dong_cd']} · 분석 기준 {('카드 주 '+scope['selected_week']) if scope.get('selected_week') else scope['ym']} · 대상 {', '.join(map(group_label,scope['target_groups']))}"),
        (0,f"자료 모드 {body['data_mode']} · {MODES.get(body['generation_mode'],body['generation_mode'])} · "+('합성 미래 시나리오·실제 예측 아님' if scope.get('synthetic_future') else '기준 자료 분석')),
        (2,'지역 및 통계 근거')]
    if body.get('proposal_version')=='regional-policy-2':
        decisions={'reuse_existing':'기존 정책 활용','adapt_existing':'기존 정책과 운영 보완','new_design_needed':'조사 범위 내 신규 설계 검토','needs_evidence':'추가 근거 확인'}
        result.append((0,{'reuse_existing':'기존 정책 활용·안내','draft':'보완·신규 대안 3개 검토'}.get(body.get('status'),'추가 근거 확인 필요')))
        if body.get('fallback_reason'):result.append((0,response_reason(body['fallback_reason'])))
        result.append((2,'필요별 대응 판단과 기존 정책 안내'))
        for row in body.get('response_decisions',[]):
            result.append((0,group_label(row['target_group'])+' · '+NEEDS.get(row['need'],'필요 가설')+' · '+decisions[row['decision']]))
            result.append((0,response_reason(row['reason'])))
            for citation in row.get('citations',[]):
                result.append((0,'판단에 사용한 공식 내용: '+citation['policy_id']+' / '+citation['field']+' / '+citation['quote']))
            for policy in row.get('guidance',[]):
                for key,label in [('title','정책'),('support','지원 내용'),('eligibility','대상'),('agency','담당 기관'),('application','신청 방법'),('deadline','기한'),('contact','문의처'),('url','출처')]:
                    result.append((0,label+': '+str(policy.get(key) or '확인 필요')))
            if row.get('outreach_draft'):
                draft=row['outreach_draft']
                result.extend([(3,'담당자 검토용 안내 초안'),(0,'제안 채널: '+str(draft.get('channel_proposal') or '담당자 확인 필요')),(0,draft['text']),(0,'실제 발송·게시하지 않은 초안입니다.')])
            result.extend((0,'보완할 요건: '+v) for v in row.get('unmet_requirements',[]))
            result.extend((0,'미확인: '+response_reason(v)) for v in row['missing_evidence'])
        result.extend((0,'설계 전 확인할 사항: '+q) for q in body['evidence'].get('design_questions',[]))
        if not selected:result.append((0,'새 정책 대안을 작성하지 않았습니다. 기존 정책 활용안 또는 추가 확인 사항을 검토하세요.'))
        profile=body['evidence'].get('regional_profile')
        if profile:
            from backend.context.contracts import METRICS
            result.append((2,'고정된 지역 인구·가구·주거 근거'))
            for fact in profile['facts']:
                result.extend([(0,METRICS[fact['metric']][0]+': '+str(fact['value'] if fact['status']=='observed' else '미확보')+' '+fact['unit']),
                    (0,f"기준 {fact['reference_year']}년 · {fact['geography_level']} · 모집단 {fact['population']} · 연령 {fact['age_definition']} · 성별 {fact['sex']}"),
                    (0,f"지역 코드 {fact['geography_code']} · 경계 {fact['geography_version']} · 표 {fact.get('source_table')} · 항목 {fact.get('source_items',[])} · 분류 {fact.get('source_dimensions',[])} · 발표 {fact.get('published_at') or '미확인'} · 원천 해시 {fact['source_hash']} · 조회 {fact['fetched_at']} · 근거 {fact['fact_id']}"),
                    (0,"관측 상태: "+fact['status']+" · 미확보 사유: "+str(fact.get('reason') or '해당 없음')),(0,fact['source_url'])])
            result.extend((0,v) for v in profile['limitations'])
        if body.get('missing_groups'):result.append((0,'누락 집단: '+', '.join(map(group_label,body['missing_groups']))))
    for fact in body['evidence']['facts']:
        if fact.get('methodology_version')=='regional-workflow-2':
            card_rows=fact.get('card_evidence') or []
            card_detail='; '.join(f"{row['source_age_band']} V {ratio(row['v_ratio'])}/{row['v_stage']}단계, N {ratio(row['n_ratio'])}/{row['n_stage']}단계" for row in card_rows) or '세부 신호 없음'
            source_stage=f"{fact.get('source_stage_label') or '원천 단계 확인 필요'} ({fact.get('source_stage') if fact.get('source_stage') is not None else '미기록'})"
            final_stage=f"{fact.get('final_stage_label') or '미산출'} ({fact.get('final_stage') if fact.get('final_stage') is not None else '미산출'})"
            partial=any(row['baseline_status']=='observed' for row in card_rows) and any(row['baseline_status']=='unavailable' for row in card_rows)
            card_status=ALARMS.get(fact['card_alarm_status'],fact['card_alarm_status'])
            if partial:card_status='부분 기준선 · '+card_status
            evidence=(f"{group_label(fact['target_group'])} · 원천 연령 {fact['source_age_band']} · 표시 틀 {fact.get('frame') or '설명용'} · 원천 단계 {source_stage} · 기본 단계 {fact.get('base_stage') if fact.get('base_stage') is not None else '산출 안 함'} · "
                      f"카드 경보 {card_status} ({fact['card_week']}, 원천 {', '.join(fact['card_source_age_bands'] or [])}; {card_detail}) · 최종 단계 {final_stage} · 조정 {fact.get('stage_adjustment', 0)}단계 · "
                      f"동별 확인 순서 {fact.get('priority_rank') or '미기록'} · 상태 {ASSESSMENTS.get(fact['assessment_status'],fact['assessment_status'])} · 연속 지수 미산출 · 근거 {fact['evidence_id']}")
        else:
            evidence=f"{fact['target_group']} · {fact['index_name']} · 값 {fact['value'] if fact['value'] is not None else '결측'} · 단계 {fact['stage'] if fact['stage'] is not None else '보류'} · 원천 기간 {fact['source_period']} · 근거 {fact['evidence_id']}"
        result.append((0,evidence))
    if not body['evidence']['facts']:result.append((0,'지역 관측 근거가 없어 현장 필요 확인이 우선입니다.'))
    for stat in body['evidence'].get('statistics',[]):
        comparison=stat.get('comparison',{})
        result.append((0,f"비교 근거 {stat.get('evidence_id')} · 비교 동 {comparison.get('peer_count',0)} · 위험 방향 차이 {comparison.get('risk_direction_difference')} · 인과 효과·위험 확률이 아닙니다."))
    result.append((2,'기존 정책·자원과 미확인 범위'))
    for row in body['evidence'].get('matching',[]):result.append((0,f"필요: {', '.join(row.get('needs',[]))} / 미충족: {', '.join(row.get('uncovered_needs',[]))} / 조건부 자원: {', '.join(row.get('selected_resource_ids',[]))}"))
    result.append((0,'검색된 후보는 전체 정책 목록이나 현지 신청 자격을 확정하지 않습니다. 예산·정원·효과는 미확인입니다.'))
    for i,a in enumerate(selected,1):
        result.extend([(2,f"{i}안. {a['title']}"),(0,f"{MODES.get(a['generation_mode'],a['generation_mode'])} · {EVIDENCE[a['evidence_status']]} · {STATES[a['review_status']]}"),
            (0,'대상: '+', '.join(map(group_label,a['target_groups']))),(0,'제안 근거: '+a['rationale']),(0,'운영 방안: '+a['proposed_service'])])
        for identifier in a.get('need_ids',[]):
            need=next((r for r in body.get('response_decisions',[]) if r['need_id']==identifier),None)
            result.append((0,'다루는 필요: '+(group_label(need['target_group'])+' · '+NEEDS.get(need['need'],'필요 가설') if need else '저장 근거 확인 필요')+' · '+identifier))
        for title,key in [('실행 순서','operating_steps'),('제안하는 담당 역할','required_roles'),('확인할 예산','budget_questions'),('형평성 위험','equity_risks'),('미확인 사항','missing_evidence')]:
            result.append((3,title));result.extend((0,'• '+line) for line in a[key])
        result.extend([(3,'평가 계획'),(0,a['evaluation_plan'])])
        if a.get('design_checks'):
            result.append((3,'지역 맞춤 설계 검토'))
            result.append((0,'구조화 초안·효과 미검증' if a['design_checks']['status']=='structured_draft' else '설계 근거 보완 필요'))
            result.extend((0,line) for line in a['design_checks']['gaps'])
        for link in a.get('context_links',[]):result.extend([(0,'지역 근거: '+link['fact_id']),(0,'설계 선택: '+link['design_choice']),(0,'확인 조건: '+link['verification_needed'])])
        if a.get('delivery_model'):result.append((0,'운영 접점: '+a['delivery_model']))
        for element in a.get('source_elements',[]):
            result.append((0,'차용할 원 정책 요소: '+element['policy_id']+' / '+element['field']+' / '+element['quote']))
        for adaptation in a.get('adaptations',[]):
            result.extend([(0,'원 정책: '+', '.join(adaptation['source_policy_ids'])),(0,'지역 차이: '+adaptation['local_difference']),
                           (0,'현지 변경안: '+adaptation['proposed_change']),(0,'검증할 조건: '+' / '.join(adaptation['verification_needed']))])
        if a.get('evaluation_design'):
            for key,label in [('outcome','관찰할 결과'),('baseline_needed','필요 기준선'),('comparison_method','비교 방법'),('observation_period','제안 관측 기간'),('success_criterion_to_agree','합의할 성공 기준'),('stop_criterion_to_agree','합의할 중단 기준')]:
                result.append((0,label+': '+a['evaluation_design'][key]))
        if a.get('operator_notes'):result.extend([(3,'담당자 수정·보완'),(0,a['operator_notes'])])
        result.append((3,'고정 출처와 조회일'))
        for c in a['citations']:
            result.extend([(0,c['title']+' · '+c['policy_id']),(0,c['evidence']['value']),(0,'조회일 '+str(c.get('checked_at') or '미확인')+' · '+('다른 관할 참고 사례·현지 자격 미확인' if c['role']=='analogue_not_local_eligibility' else '조회된 정책 후보'))])
            if str(c.get('url','')).startswith(('https://','http://')):result.append((0,c['url']))
        if not a['citations']:result.append((0,'인용할 공식 정책 근거 없음. 근거 보완용 구상안입니다.'))
    result.append((2,'검토 이력과 한계'))
    for h in body.get('history',[]):result.append((0,f"revision {h['revision']} · {h['operator']} · {h['action']} · {h['text']}"))
    result.append((0,'지역 집계 지표는 개인 진단이 아닙니다. 제안의 타당성·실행 가능성·효과는 담당자와 현장 검증이 필요합니다.'))
    return result


def render_pdf(content):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate,Paragraph
    if 'NanumGothic' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('NanumGothic',str(FONT)))
    styles={level:ParagraphStyle(str(level),fontName='NanumGothic',fontSize={0:10,1:20,2:14,3:11}[level],leading={0:16,1:28,2:22,3:18}[level],wordWrap='CJK',spaceBefore=8 if level else 0,spaceAfter=6,keepWithNext=bool(level)) for level in range(4)}
    stream=BytesIO();story=[Paragraph(escape(text).replace('\n','<br/>'),styles[level]) for level,text in content]
    def footer(canvas,doc):
        canvas.saveState();canvas.setFont('NanumGothic',8);canvas.drawString(42,24,'검토용 초안 · 실제 배정 아님');canvas.drawRightString(A4[0]-42,24,str(doc.page));canvas.restoreState()
    SimpleDocTemplate(stream,pagesize=A4,leftMargin=42,rightMargin=42,topMargin=42,bottomMargin=42,title=content[0][1]).build(story,onFirstPage=footer,onLaterPages=footer)
    return stream.getvalue()


def render_docx(content):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    document=Document()
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3'):
        style=document.styles[name];style.font.name='NanumGothic';style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'NanumGothic')
    document.styles['Normal'].font.size=Pt(10)
    for level,text in content:
        if level:document.add_heading(text,level=level-1)
        else:document.add_paragraph(text)
    document.sections[0].footer.paragraphs[0].text='검토용 초안 · 실제 배정 아님'
    stream=BytesIO();document.save(stream);return stream.getvalue()


def router(store):
    result=APIRouter(prefix='/policy-proposals')
    @result.get('/{owner_id}/{dong_cd}/download/{format}')
    def download(owner_id:UUID,dong_cd:str,format:Literal['pdf','docx'],version:str,revision:int=Query(ge=0),alternative_id:UUID|None=None):
        try:
            with store.connection() as db:body=read(db,str(owner_id),dong_cd,revision)
            if body['version']!=version:raise HTTPException(409,'제안 버전이 일치하지 않습니다.')
            content=sections(body,str(alternative_id) if alternative_id else None)
        except KeyError:raise HTTPException(404,'요청한 제안·대안·검토 버전 없음')
        raw=render_pdf(content) if format=='pdf' else render_docx(content)
        filename=f'policy-{owner_id}-r{revision}.{format}'
        return Response(raw,media_type='application/pdf' if format=='pdf' else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':f'attachment; filename="{filename}"','Cache-Control':'no-store'})
    return result
